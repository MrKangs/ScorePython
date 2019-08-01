# import openpyxl module 
import openpyxl 
  
# import ScatterChart, Reference, Series 
# class from openpyxl.chart sub_module 
from openpyxl.chart import ScatterChart, Reference, Series 
  
# Call a Workbook() function of openpyxl  
# to create a new blank Workbook object 
wb = openpyxl.Workbook() 
  
# Get workbook active sheet  
# from the active attribute. 
sheet = wb.active 
  
rows = [ 
    ("Number of Products", "Sales in USD", "Market share"), 
    (14, 12200, 15), 
    (20, 60000, 33), 
    (18, 24400, 10), 
    (22, 32000, 42), 
] 
  
# write content of each row in 1st, 2nd and 3rd 
# column of the active sheet respectively . 
for row in rows: 
    sheet.append(row) 
  
# Create object of ScatterChart class 
chart = ScatterChart() 
  
# create data for plotting 
xvalues = Reference(sheet, min_col = 1, 
                    min_row = 2, max_row = 5) 
                      
yvalues = Reference(sheet, min_col = 2, 
                    min_row = 2, max_row = 5) 
                      
size = Reference(sheet, min_col = 3, 
                 min_row = 2, max_row = 5) 
  
# create a 1st series of data 
series = Series(values = yvalues, xvalues = xvalues, 
                      zvalues = size, title ="2013") 
  
# add series data to the chart object 
chart.series.append(series) 
  
# set the title of the chart 
chart.title = " SCATTER-CHART "
  
# set the title of the x-axis 
chart.x_axis.title = " X_AXIS "
  
# set the title of the y-axis 
chart.y_axis.title = " Y_AXIS "
  
# add chart to the sheet 
# the top-left corner of a chart 
# is anchored to cell E2 . 
sheet.add_chart(chart, "E2") 
  
# save the file 
wb.save(" ScatterChart.xlsx") 